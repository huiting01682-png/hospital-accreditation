"""由 data/standards.csv 與 data/units.csv 產生評鑑條文對照表 Excel。

用法：python3 tools/build_xlsx.py [輸出路徑]
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "output" / "accreditation_checklist.xlsx"

# 條文本身的欄位（來自條文檔）
STANDARD_COLS = ["篇", "章", "條號", "條文名稱", "條文類別", "評量項目"]
# 各單位追蹤填寫的欄位
TRACK_COLS = ["負責單位", "聯絡人", "佐證資料", "自評等級", "進度",
              "預計完成日", "委員建議事項", "待改善事項", "更新日期"]
COLS = STANDARD_COLS + TRACK_COLS
WIDTHS = {"篇": 12, "章": 18, "條號": 10, "條文名稱": 30, "條文類別": 10, "評量項目": 50,
          "負責單位": 14, "聯絡人": 10, "佐證資料": 30, "自評等級": 10, "進度": 10,
          "預計完成日": 12, "委員建議事項": 36, "待改善事項": 30, "更新日期": 12}

GRADES = ["優良", "符合", "待改善", "不適用"]
PROGRESS = ["未開始", "進行中", "已完成", "需協助"]
CATEGORIES = ["必要", "重點", "可免評", "試評", "一般"]

INK, NAVY, SLATE = "1B2559", "2E5B8F", "5A6478"
MINT, AMBER, CORAL = "0E8A63", "D97706", "E63946"
FONT = "Microsoft JhengHei"
thin = Side(style="thin", color="C9D1DC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=12)
BODY_FONT = Font(name=FONT, size=12)
WRAP = Alignment(wrap_text=True, vertical="top")


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def header(ws, cols, fill):
    for i, name in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font, c.fill, c.border = HEAD_FONT, PatternFill("solid", fgColor=fill), BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def build():
    standards = read_csv(ROOT / "data" / "standards.csv")
    units = [r["負責單位"] for r in read_csv(ROOT / "data" / "units.csv") if r["負責單位"].strip()]
    nrows = max(len(standards), 1) + 300  # 預留空白列供手動新增
    last = nrows + 1

    wb = Workbook()

    # ── 條文總表 ──
    ws = wb.active
    ws.title = "條文總表"
    header(ws, COLS, NAVY)
    for c in range(len(STANDARD_COLS) + 1, len(COLS) + 1):  # 追蹤欄位用 SOFT 色區分
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="3F4EA8")
    for r, row in enumerate(standards, 2):
        for c, name in enumerate(COLS, 1):
            ws.cell(row=r, column=c, value=row.get(name, "") or None)
    for r in range(2, last + 1):
        for c, name in enumerate(COLS, 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.alignment, cell.border = BODY_FONT, WRAP, BORDER
            if name in ("預計完成日", "更新日期"):
                cell.number_format = "yyyy/mm/dd"
    for i, name in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[name]
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{last}"

    def col(name):
        return get_column_letter(COLS.index(name) + 1)

    def dropdown(name, source):
        dv = DataValidation(type="list", formula1=source, allow_blank=True)
        dv.error, dv.errorTitle = "請從下拉清單選擇", "輸入值不在清單內"
        ws.add_data_validation(dv)
        dv.add(f"{col(name)}2:{col(name)}{last}")

    dropdown("負責單位", f"=單位清單!$A$2:$A${len(units) + 50}")
    dropdown("自評等級", '"' + ",".join(GRADES) + '"')
    dropdown("進度", '"' + ",".join(PROGRESS) + '"')
    dropdown("條文類別", '"' + ",".join(CATEGORIES) + '"')
    date_dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
    date_dv.error = "請輸入日期，例如 2026/12/31"
    ws.add_data_validation(date_dv)
    for name in ("預計完成日", "更新日期"):
        date_dv.add(f"{col(name)}2:{col(name)}{last}")

    # 燈號：進度與自評等級
    def light(name, value, color):
        rng = f"{col(name)}2:{col(name)}{last}"
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=[f'"{value}"'],
            font=Font(name=FONT, bold=True, color=color)))
    light("進度", "已完成", MINT)
    light("進度", "進行中", AMBER)
    light("進度", "需協助", CORAL)
    light("自評等級", "待改善", CORAL)
    # 預計完成日已過但未完成 → 期程差異
    due, prog = col("預計完成日"), col("進度")
    from openpyxl.formatting.rule import FormulaRule
    ws.conditional_formatting.add(f"{due}2:{due}{last}", FormulaRule(
        formula=[f'AND({due}2<>"",{due}2<TODAY(),{prog}2<>"已完成")'],
        font=Font(name=FONT, bold=True, color=CORAL)))

    # ── 單位清單 ──
    wu = wb.create_sheet("單位清單")
    header(wu, ["負責單位"], NAVY)
    for r, u in enumerate(units, 2):
        wu.cell(row=r, column=1, value=u).font = BODY_FONT
    wu.column_dimensions["A"].width = 20
    wu.cell(row=1, column=3, value="可直接在 A 欄新增或修改單位名稱，條文總表的下拉選單會同步更新。").font = Font(name=FONT, color=SLATE)

    # ── 進度統計 ──
    st = wb.create_sheet("進度統計")
    T = "條文總表!"
    rng = lambda name: f"{T}${col(name)}$2:${col(name)}${last}"
    cols = ["負責單位", "條文數"] + PROGRESS + ["完成率", "待改善", "期程差異"]
    header(st, cols, INK)
    for r, u in enumerate(units, 2):
        st.cell(row=r, column=1, value=f"=單位清單!A{r}")
        st.cell(row=r, column=2, value=f'=COUNTIF({rng("負責單位")},A{r})')
        for i, p in enumerate(PROGRESS):
            st.cell(row=r, column=3 + i, value=f'=COUNTIFS({rng("負責單位")},$A{r},{rng("進度")},"{p}")')
        st.cell(row=r, column=7, value=f'=IF(B{r}=0,"",E{r}/B{r})').number_format = "0%"
        st.cell(row=r, column=8, value=f'=COUNTIFS({rng("負責單位")},$A{r},{rng("自評等級")},"待改善")')
        st.cell(row=r, column=9, value=(
            f'=COUNTIFS({rng("負責單位")},$A{r},{rng("預計完成日")},"<"&TODAY(),{rng("進度")},"<>已完成")'))
    tr = len(units) + 2
    st.cell(row=tr, column=1, value="全院合計")
    st.cell(row=tr, column=2, value=f'=COUNTA({rng("條號")})')
    for i, p in enumerate(PROGRESS):
        st.cell(row=tr, column=3 + i, value=f'=COUNTIF({rng("進度")},"{p}")')
    st.cell(row=tr, column=7, value=f'=IF(B{tr}=0,"",E{tr}/B{tr})').number_format = "0%"
    st.cell(row=tr, column=8, value=f'=COUNTIF({rng("自評等級")},"待改善")')
    st.cell(row=tr, column=9, value=f'=COUNTIFS({rng("預計完成日")},"<"&TODAY(),{rng("進度")},"<>已完成")')
    st.cell(row=tr + 1, column=1, value="尚未指派單位")
    st.cell(row=tr + 1, column=2, value=f'=COUNTIFS({rng("條號")},"<>",{rng("負責單位")},"")')
    for r in range(2, tr + 2):
        for c in range(1, len(cols) + 1):
            cell = st.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = Font(name=FONT, size=12, bold=r >= tr)
    for i, w in enumerate([16, 10, 10, 10, 10, 10, 10, 10, 10], 1):
        st.column_dimensions[get_column_letter(i)].width = w
    st.cell(row=tr + 3, column=1, value="期程差異：預計完成日已過、進度尚未標示為已完成的條數。").font = Font(name=FONT, color=SLATE)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"已輸出 {OUT}（條文 {len(standards)} 條，單位 {len(units)} 個）")


if __name__ == "__main__":
    build()
